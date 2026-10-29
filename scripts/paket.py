"""Paket za razrednika: Word i PowerPoint materijali koje agencija šalje uz ponudu.

Pokretanje (iz korijena repozitorija):
    node scripts/export-trips.mjs > /tmp/trips.json && python3 scripts/paket.py /tmp/trips.json

Rezultat: materijali/paket-za-razrednika/ (opći dokumenti + mapa po programu).
Uglate zagrade [..] su mjesta koja razrednik ili agencija popunjava prije slanja.
Potrebno: pip install python-docx python-pptx
"""
import json
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor, Cm
from pptx import Presentation
from pptx.dml.color import RGBColor as PRGB
from pptx.util import Inches, Pt as PPt

BLUE = RGBColor(0x10, 0x3A, 0x6E)
AMBER = RGBColor(0xD0, 0x74, 0x10)
GREY = RGBColor(0x55, 0x60, 0x70)
OUT = Path('materijali/paket-za-razrednika')

# Ciljevi i ishodi po programu (prijedlog za školski kurikulum; razrednik ih prilagođava)
GOALS = {
    'skola-u-prirodi-gorski-kotar': {
        'cilj': 'Upoznati prirodno-geografska obilježja gorske Hrvatske, biljni i životinjski svijet planinske šume te život ljudi u gorskom kraju; razvijati samostalnost i suradnju u skupini.',
        'ishodi': [
            'opisuje obilježja gorskog zavičaja (reljef, klima, vode, biljke i životinje)',
            'prepoznaje tragove životinja i razlikuje biljke planinske šume',
            'snalazi se u prostoru pomoću karte i strana svijeta',
            'objašnjava važnost zaštite prirode na primjeru nacionalnog parka',
            'sudjeluje u zajedničkim aktivnostima i preuzima odgovornost za sebe i skupinu',
        ],
        'vrednovanje': 'Plakat i izlaganje skupine, kviz znanja, likovni radovi, razgovor nakon povratka.',
    },
    'zagrebacke-price': {
        'cilj': 'Upoznati glavni grad Republike Hrvatske, njegovu povijesnu jezgru i važne ustanove te povezati baštinu s tehnikom i znanošću.',
        'ishodi': [
            'opisuje znamenitosti Gornjeg grada i Kaptola te njihovu ulogu u povijesti grada',
            'objašnjava ulogu glavnog grada i važnih državnih ustanova',
            'povezuje izloške Tehničkog muzeja s primjenom u svakodnevnom životu',
            'prepričava legendu ili priču vezanu uz zagrebačku znamenitost',
        ],
        'vrednovanje': 'Radni listić, kratki sastavak ili prezentacija o jednoj znamenitosti.',
    },
    'istra-brijuni': {
        'cilj': 'Upoznati prirodnu i kulturnu baštinu Istre od antike do danas te zaštićena područja i njihovu ulogu.',
        'ishodi': [
            'opisuje rimsku baštinu Pule i njezino značenje',
            'objašnjava razloge zaštite otočja Brijuni i razlikuje autohtone i unesene vrste',
            'uspoređuje obalna i unutrašnja naselja Istre',
            'prikuplja i prezentira podatke o jednoj lokaciji',
        ],
        'vrednovanje': 'Projektni zadatak skupine, prezentacija ili plakat, kviz.',
    },
    'lika-plitvice': {
        'cilj': 'Upoznati krške pojave, nastanak sedrenih barijera i vrijednosti nacionalnog parka te život i djelo Nikole Tesle.',
        'ishodi': [
            'objašnjava nastanak sedre i jezera te obilježja krša',
            'opisuje korištenje vodene snage na primjeru mlinova u Rastokama',
            'navodi najvažnija Teslina otkrića i njihovu primjenu',
            'objašnjava pravila ponašanja u zaštićenom području',
        ],
        'vrednovanje': 'Radni listić, fotografski zadatak, razgovor nakon povratka.',
    },
    'bec-budimpesta': {
        'cilj': 'Upoznati povijest srednje Europe i Habsburške Monarhije, kulturnu baštinu Beča i Budimpešte te povezanost Hrvatske s tim prostorom.',
        'ishodi': [
            'objašnjava ulogu Beča i Budimpešte u povijesti Habsburške Monarhije i Hrvatske',
            'opisuje odabrane umjetničke i arhitektonske stilove na primjerima',
            'koristi strani jezik u svakodnevnim situacijama',
            'uspoređuje dva grada na Dunavu prema zadanim kriterijima',
        ],
        'vrednovanje': 'Prezentacija skupine, putopis ili kviz.',
    },
    'varazdin-trakoscan': {
        'cilj': 'Upoznati baroknu baštinu Varaždina, dvorac Trakošćan i prirodne vrijednosti sjeverozapadne Hrvatske.',
        'ishodi': [
            'opisuje obilježja baroka na primjerima varaždinske jezgre',
            'objašnjava ulogu dvoraca i plemićkih obitelji u povijesti kraja',
            'razlikuje skupine kukaca na primjerima iz entomološke zbirke',
        ],
        'vrednovanje': 'Radni listić, crtež ili kratka prezentacija.',
    },
}

AGES = {'razredna': '1. – 4. razred', 'predmetna': '5. – 8. razred', 'srednja': 'srednja škola'}


def duration(t):
    d, n = t['days'], t['nights']
    day = 'dan' if d == 1 else 'dana'
    if not n:
        return f'{d} {day}'
    night = 'noć' if n == 1 else 'noći'
    return f'{d} {day} / {n} {night}'


def base_doc():
    doc = Document()
    st = doc.styles['Normal']
    st.font.name = 'Calibri'
    st.font.size = Pt(11)
    for s in doc.sections:
        s.left_margin = s.right_margin = Cm(2.2)
        s.top_margin = s.bottom_margin = Cm(2)
    return doc


def header(doc, data, title, subtitle=None):
    p = doc.add_paragraph()
    r = p.add_run('Veliki odmor · školski izleti i ekskurzije')
    r.bold = True
    r.font.color.rgb = BLUE
    r.font.size = Pt(10)
    o = data['owner']
    p2 = doc.add_paragraph()
    r2 = p2.add_run(f"{o['fullName']}, {o['street']}, {o['city']} · OIB {o['oib']} · ID kod {o['idCode']} · {data['emails'][0]} · {data['phones'][0]}")
    r2.font.size = Pt(8)
    r2.font.color.rgb = GREY
    h = doc.add_heading(title, level=1)
    for run in h.runs:
        run.font.color.rgb = BLUE
    if subtitle:
        sp = doc.add_paragraph()
        sr = sp.add_run(subtitle)
        sr.italic = False
        sr.font.color.rgb = GREY


def h2(doc, text):
    h = doc.add_heading(text, level=2)
    for run in h.runs:
        run.font.color.rgb = BLUE


def table(doc, rows, widths=(5.5, 11)):
    tb = doc.add_table(rows=0, cols=2)
    tb.style = 'Table Grid'
    for k, v in rows:
        cells = tb.add_row().cells
        cells[0].text = k
        cells[1].text = v
        cells[0].paragraphs[0].runs[0].bold = True
        cells[0].width = Cm(widths[0])
        cells[1].width = Cm(widths[1])
    doc.add_paragraph()
    return tb


def bullets(doc, items):
    for i in items:
        doc.add_paragraph(i, style='List Bullet')


# 1. Ciljevi i ishodi za školski kurikulum (po programu)
def kurikulum(data, t, folder):
    g = GOALS[t['id']]
    doc = base_doc()
    header(doc, data, f"Prijedlog za školski kurikulum: {t['title']}", 'Prijedlog za prilagodbu; konačni tekst određuje škola.')
    table(doc, [
        ('Aktivnost', f"{'Jednodnevni izlet' if t['type'] == 'jednodnevni' else 'Višednevna izvanučionička nastava'}: {t['title']}"),
        ('Razred', ', '.join(AGES[a] for a in t['ages']) + ' [razred i odjeljenja]'),
        ('Odredište', t['region']),
        ('Trajanje', duration(t)),
        ('Cilj', g['cilj']),
        ('Povezanost s predmetima', ', '.join(t['subjects'])),
        ('Nositelji', '[Razrednici / učitelji pratitelji], turistička agencija TERRA POMALO d.o.o. (Veliki odmor)'),
        ('Način realizacije', 'Terenska nastava prema programu putovanja uz stručno vodstvo; radni listići i zadaci za učenike.'),
        ('Vremenik', '[Mjesec i godina]'),
        ('Troškovnik', 'Prema odabranoj ponudi na javni poziv; troškove snose roditelji [moguće obročno plaćanje].'),
        ('Način vrednovanja', g['vrednovanje']),
    ])
    h2(doc, 'Očekivani ishodi')
    doc.add_paragraph('Učenik:')
    bullets(doc, g['ishodi'])
    doc.save(folder / '1-prijedlog-za-kurikulum.docx')


# 2. Sažetak programa na jednoj stranici (za roditeljski sastanak)
def sazetak(data, t, folder):
    doc = base_doc()
    header(doc, data, t['title'], f"{t['region']} · {duration(t)} · {', '.join(AGES[a] for a in t['ages'])}")
    doc.add_paragraph(t['summary'])
    h2(doc, 'Program')
    for d in t['itinerary']:
        p = doc.add_paragraph()
        r = p.add_run(d['label'] + ': ')
        r.bold = True
        p.add_run('; '.join(text for _, text in d['items']))
    h2(doc, 'Cijena uključuje')
    bullets(doc, t['included'])
    h2(doc, 'Cijena i plaćanje')
    doc.add_paragraph('[… € po učeniku] za [broj] učenika. Plaćanje jednokratno, karticama ili u ratama: [plan rata].')
    doc.add_paragraph('Osiguranje od nesretnog slučaja uključeno je u cijenu; osiguranje od otkaza putovanja je izborno.')
    h2(doc, 'Kontakt')
    doc.add_paragraph(f"Koordinator tijekom putovanja (0–24 h): {data['phones'][0]} · {data['emails'][0]}")
    doc.save(folder / '2-sazetak-programa.docx')


# 3. Prezentacija za roditeljski sastanak (do 5 minuta)
def prezentacija(data, t, folder):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    blank = prs.slide_layouts[6]

    def slide(title, lines, sub=None):
        s = prs.slides.add_slide(blank)
        bg = s.background.fill
        bg.solid()
        bg.fore_color.rgb = PRGB(0xF7, 0xFB, 0xFE)
        tb = s.shapes.add_textbox(Inches(0.8), Inches(0.6), Inches(11.7), Inches(1.2)).text_frame
        tb.word_wrap = True
        r = tb.paragraphs[0].add_run()
        r.text = title
        r.font.size = PPt(36)
        r.font.bold = True
        r.font.color.rgb = PRGB(0x10, 0x3A, 0x6E)
        if sub:
            p = tb.add_paragraph()
            rr = p.add_run()
            rr.text = sub
            rr.font.size = PPt(18)
            rr.font.color.rgb = PRGB(0x55, 0x60, 0x70)
        body = s.shapes.add_textbox(Inches(0.8), Inches(2.1), Inches(11.7), Inches(4.8)).text_frame
        body.word_wrap = True
        for i, line in enumerate(lines):
            p = body.paragraphs[0] if i == 0 else body.add_paragraph()
            rr = p.add_run()
            rr.text = '• ' + line
            rr.font.size = PPt(20)
            rr.font.color.rgb = PRGB(0x08, 0x1F, 0x3D)
            p.space_after = PPt(10)
        foot = s.shapes.add_textbox(Inches(0.8), Inches(6.9), Inches(11.7), Inches(0.4)).text_frame
        fr = foot.paragraphs[0].add_run()
        fr.text = f"Veliki odmor · {data['owner']['name']} · ID kod {data['owner']['idCode']}"
        fr.font.size = PPt(11)
        fr.font.color.rgb = PRGB(0xD0, 0x74, 0x10)

    slide(t['title'], [t['summary'], f"Trajanje: {duration(t)}", f"Povezanost s predmetima: {', '.join(t['subjects'])}"], t['region'])
    for d in t['itinerary']:
        slide(d['label'], [f'{time}  {text}' for time, text in d['items']][:7], t['title'])
    slide('Cijena uključuje', t['included'][:7])
    slide('Sigurnost', [
        'Prijevoznici registrirani za prijevoz putnika, turistički autobusi',
        'Smještaj provjeren za školske grupe; noćni nadzor dogovoren sa školom',
        'Alergije i posebnu prehranu usklađujemo prije polaska',
        'Osiguranje od nesretnog slučaja u cijeni',
        f"Koordinator 0–24 h: {data['phones'][0]}",
    ])
    slide('Cijena i plaćanje', ['[… € po učeniku] za [broj] učenika', 'Plaćanje jednokratno, karticama ili u ratama: [plan rata]', 'Gratis za pratitelje prema ugovoru sa školom', 'Osiguranje od otkaza putovanja je izborno'])
    slide('Pitanja?', [f"{data['emails'][0]} · {data['phones'][0]}", 'Odgovore za roditelje i Opće uvjete pogledajte na našoj stranici [adresa stranice]'])
    prs.save(folder / '3-prezentacija-roditeljski-sastanak.pptx')


# Opći dokumenti (isti za sve programe)
def obavijest(data):
    doc = base_doc()
    header(doc, data, 'Obavijest roditeljima i suglasnost za putovanje')
    doc.add_paragraph('Poštovani roditelji,')
    doc.add_paragraph('obavještavamo vas da [naziv škole] organizira [izlet / ekskurziju / školu u prirodi] za učenike [razred] razreda. '
                      'Putovanje je dio školskog kurikuluma, a organizator je turistička agencija TERRA POMALO d.o.o. (Veliki odmor).')
    table(doc, [
        ('Odredište i program', '[naziv programa, ukratko]'),
        ('Termin', '[datum polaska] – [datum povratka]'),
        ('Polazak i povratak', '[mjesto i okvirno vrijeme]'),
        ('Voditelj putovanja (škola)', '[ime i prezime, kontakt]'),
        ('Cijena po učeniku', '[… €]; plaćanje: [jednokratno / u … rata, rokovi]'),
        ('Uplata', f"{data['owner']['name']}, IBAN {data['owner']['iban']}, poziv na broj: [broj ugovora i ime učenika]"),
        ('Roditeljski sastanak', '[datum i vrijeme]'),
    ])
    h2(doc, 'Suglasnost roditelja ili skrbnika')
    doc.add_paragraph('Ja, ______________________________ (ime i prezime roditelja/skrbnika), suglasan/suglasna sam da moje dijete '
                      '______________________________ (ime i prezime učenika), razred ______, sudjeluje na putovanju [naziv programa] u terminu [termin].')
    doc.add_paragraph('Upoznat/a sam s programom putovanja, cijenom, načinom plaćanja i Općim uvjetima organizatora.')
    h2(doc, 'Podaci važni za putovanje')
    table(doc, [
        ('Alergije (hrana, lijekovi, ostalo)', ''),
        ('Posebna prehrana', ''),
        ('Lijekovi koje dijete uzima i upute', ''),
        ('Druge napomene za pratitelje', ''),
        ('Kontakt roditelja tijekom putovanja (telefon)', ''),
    ])
    doc.add_paragraph('Podatke o zdravlju i prehrani dajem izričitom privolom, isključivo radi sigurnosti djeteta na ovom putovanju; '
                      'škola i organizator ih ne koriste u druge svrhe i brišu nakon putovanja.')
    doc.add_paragraph('\nMjesto i datum: ____________________        Potpis roditelja/skrbnika: ____________________')
    doc.save(OUT / 'obavijest-i-suglasnost-roditelja.docx')


def pakiranje(data):
    doc = base_doc()
    header(doc, data, 'Popis za pakiranje i pravila ponašanja', 'Za učenike i roditelje; škola ga može dopuniti.')
    h2(doc, 'Za sva putovanja')
    bullets(doc, ['Zdravstvena iskaznica (za inozemstvo i osobna iskaznica ili putovnica te europska kartica zdravstvenog osiguranja)',
                  'Lijekovi koje dijete redovito uzima, s uputom za pratitelje', 'Boca za vodu i manji ruksak za dnevne izlete',
                  'Udobna obuća za hodanje i kabanica ili kišobran', 'Punjač za mobitel (prema pravilima škole)', 'Džeparac prema dogovoru s razrednikom'])
    h2(doc, 'Dodatno za višednevna putovanja')
    bullets(doc, ['Odjeća za svaki dan i jedan rezervni komplet', 'Pidžama, papuče, pribor za osobnu higijenu i ručnik (ako ga smještaj ne osigurava)',
                  'Topla jakna ili flis i kapa (za planinska odredišta i u proljeće)', 'Sportska odjeća i obuća za aktivnosti', 'Kupaći kostim, ako program uključuje kupanje'])
    h2(doc, 'Pravila ponašanja')
    bullets(doc, ['Slušam upute pratitelja, voditelja putovanja i vodiča', 'Ne udaljavam se od grupe bez dopuštenja i držim se dogovorenog vremena',
                  'U autobusu sjedim na svom mjestu i vežem pojas', 'Poštujem mir u smještaju i ne izlazim iz sobe noću bez dopuštenja',
                  'Ne konzumiram alkohol, cigarete ni druga sredstva ovisnosti', 'Čuvam svoje stvari i tuđu imovinu; štetu nadoknađuju roditelji',
                  'Odmah javljam pratitelju ako se ne osjećam dobro ili ako se nešto dogodi'])
    doc.save(OUT / 'popis-za-pakiranje-i-pravila-ponasanja.docx')


def izvjesce(data):
    doc = base_doc()
    header(doc, data, 'Izvješće o provedenom putovanju', 'Predložak za voditelja putovanja (škola).')
    table(doc, [
        ('Škola', ''), ('Razred i broj učenika', ''), ('Pratitelji', ''), ('Program i odredište', ''),
        ('Termin', ''), ('Organizator', f"{data['owner']['name']} (Veliki odmor), ID kod {data['owner']['idCode']}"),
    ])
    h2(doc, 'Realizacija programa')
    doc.add_paragraph('[Je li program proveden prema planu; promjene i razlozi]')
    h2(doc, 'Ostvareni ciljevi i ishodi')
    doc.add_paragraph('[Ciljevi iz kurikuluma i kako su ostvareni; zadaci i radovi učenika]')
    h2(doc, 'Prijevoz, smještaj i prehrana')
    doc.add_paragraph('[Ocjena i napomene]')
    h2(doc, 'Posebni događaji')
    doc.add_paragraph('[Zdravstveni ili drugi događaji i postupanje; ako ih nije bilo, navesti „nije bilo“]')
    h2(doc, 'Zadovoljstvo učenika i roditelja')
    doc.add_paragraph('[Sažetak dojmova ili ankete]')
    h2(doc, 'Prijedlozi')
    doc.add_paragraph('[Za sljedeća putovanja]')
    doc.add_paragraph('\nMjesto i datum: ____________________        Voditelj putovanja: ____________________')
    doc.save(OUT / 'predlozak-izvjesca.docx')


def readme(data):
    names = '\n'.join(f"- `{t['id']}/` – {t['title']}" for t in data['trips'])
    (OUT / 'README.md').write_text(f"""# Paket za razrednika

Materijali koje Veliki odmor šalje razredniku uz ponudu. Generira ih `scripts/paket.py` iz podataka programa,
pa se nakon promjene programa ponovo pokrenu (upute u skripti). [Uglate zagrade] popunjava razrednik ili agencija.

Opći dokumenti:
- `obavijest-i-suglasnost-roditelja.docx` – obavijest, suglasnost i podaci o zdravlju i prehrani (izričita privola)
- `popis-za-pakiranje-i-pravila-ponasanja.docx`
- `predlozak-izvjesca.docx` – izvješće voditelja putovanja nakon povratka

Po programu (mapa s id-om programa):
- `1-prijedlog-za-kurikulum.docx` – cilj, ishodi, nositelji, vremenik, troškovnik, vrednovanje
- `2-sazetak-programa.docx` – jedna stranica za roditeljski sastanak
- `3-prezentacija-roditeljski-sastanak.pptx` – do 5 minuta

{names}

Napomena: tekst javnog poziva i uvjete ponude piše škola; agencija ih ne priprema.
""", encoding='utf-8')


def main(path):
    data = json.loads(Path(path).read_text(encoding='utf-8'))
    OUT.mkdir(parents=True, exist_ok=True)
    obavijest(data)
    pakiranje(data)
    izvjesce(data)
    for t in data['trips']:
        folder = OUT / t['id']
        folder.mkdir(exist_ok=True)
        kurikulum(data, t, folder)
        sazetak(data, t, folder)
        prezentacija(data, t, folder)
    readme(data)
    print(f'Paket je u {OUT}/')


if __name__ == '__main__':
    main(sys.argv[1])
